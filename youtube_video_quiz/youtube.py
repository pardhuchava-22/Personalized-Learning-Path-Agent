# app.py
import os
import re
import json
import subprocess
import urllib.parse
import io # NEW: For sending file from memory
from flask import Flask, request, jsonify
from flask_cors import CORS
from pathlib import Path
from youtube_transcript_api import YouTubeTranscriptApi
import requests
import google.generativeai as genai
from dotenv import load_dotenv
from docx import Document # NEW: For creating .docx files
from docx.shared import Pt, Inches

# --- Load envionment variables FIRST ---
load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash") # Using 2026 gemini-2.5-flash
ASSEMBLYAI_API_KEY = os.getenv("ASSEMBLYAI_API_KEY")

if not GEMINI_API_KEY:
    raise ValueError("GEMINI_API_KEY not found in environment variables!")

print(f"✓ Loaded API Key: {GEMINI_API_KEY[:10]}...")

# --- Setup ---
app = Flask(__name__)
CORS(app, resources={
    r"/*": {
        "origins": "*",
        "methods": ["GET", "POST", "OPTIONS"],
        "allow_headers": ["Content-Type"]
    }
})

genai.configure(api_key=GEMINI_API_KEY)

TRANSCRIPTS_DIR = Path("./transcripts")
TRANSCRIPTS_DIR.mkdir(exist_ok=True)

# --- Helper Functions (Existing) ---

def extract_youtube_id(url):
    """Extracts YouTube video ID from standard or short URLs."""
    parsed = urllib.parse.urlparse(url)
    if parsed.hostname in ['www.youtube.com', 'youtube.com']:
        q = urllib.parse.parse_qs(parsed.query)
        return q.get('v', [None])[0]
    elif parsed.hostname in ['youtu.be']:
        return parsed.path[1:]
    return None

def get_video_duration(video_id):
    """Get video duration in seconds using yt-dlp"""
    try:
        result = subprocess.run([
            "yt-dlp",
            "--print", "duration",
            f"https://www.youtube.com/watch?v={video_id}"
        ], capture_output=True, text=True, check=True)
        
        duration = float(result.stdout.strip())
        return int(duration)
    except Exception as e:
        print(f"Could not get video duration: {e}")
        return 3600

def get_transcript_with_timestamps(video_id):
    """Get transcript with timestamps."""
    transcript_json_path = TRANSCRIPTS_DIR / f"{video_id}_timestamps.json"
    
    if transcript_json_path.exists():
        with open(transcript_json_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    
    try:
        api = YouTubeTranscriptApi()
        fetched_data = api.fetch(video_id, languages=['en'])
        transcript_list = []
        for seg in fetched_data:
            if isinstance(seg, dict):
                transcript_list.append(seg)
            else:
                transcript_list.append({
                    'text': getattr(seg, 'text', ''),
                    'start': getattr(seg, 'start', 0.0),
                    'duration': getattr(seg, 'duration', 0.0)
                })
        with open(transcript_json_path, 'w', encoding='utf-8') as f:
            json.dump(transcript_list, f)
        return transcript_list
    except Exception as e:
        print(f"YouTubeTranscriptApi failed: {e}")
        return None

# --- NEW: Transcript Cleaner ---
def clean_transcript_text(text):
    """Removes common junk and sponsor messages from transcript text."""
    # Remove metadata lines
    text = re.sub(r'Kind: captions.*', '', text, flags=re.IGNORECASE)
    text = re.sub(r'Language: en.*', '', text, flags=re.IGNORECASE)
    
    # Remove common sponsor phrases
    sponsor_keywords = ['sponsor', 'dashlane', 'nordvpn', 'squarespace', 'brilliant.org', 'audible']
    # Create a regex to match sentences containing these keywords
    for keyword in sponsor_keywords:
        # Matches a full sentence (between periods) containing the keyword
        text = re.sub(r"([^.!?]*" + re.escape(keyword) + r"[^.!?]*[.!?])", "", text, flags=re.IGNORECASE)
    
    # Remove repetitive intro/outro phrases (optional but helpful)
    text = re.sub(r'(hi|hello|hey) (everybody|everyone|guys|folks)', '', text, flags=re.IGNORECASE)
    text = re.sub(r'(thanks for watching|see you in the next video)', '', text, flags=re.IGNORECASE)
    
    # Remove timestamps
    text = re.sub(r'\d{2}:\d{2}:\d{2}.\d{3} --> \d{2}:\d{2}:\d{2}.\d{3}', '', text)
    
    # Remove extra whitespace
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def get_transcript_text(video_id):
    """Get plain text transcript (fallback)"""
    transcript_path = TRANSCRIPTS_DIR / f"{video_id}.txt"
    
    if transcript_path.exists():
        return transcript_path.read_text(encoding="utf-8")

    try:
        subprocess.run([
            "yt-dlp",
            "--write-auto-subs",
            "--sub-lang", "en",
            "--skip-download",
            "-o", str(TRANSCRIPTS_DIR / f"{video_id}"),
            f"https://www.youtube.com/watch?v={video_id}"
        ], check=True, capture_output=True, text=True)

        vtt_files = list(TRANSCRIPTS_DIR.glob(f"{video_id}*.vtt"))
        if vtt_files:
            vtt_file = vtt_files[0]
            lines = vtt_file.read_text(encoding="utf-8").splitlines()
            full_text = " ".join(
                re.sub(r'<[^>]+>', '', line) 
                for line in lines 
                if line.strip() 
                and "-->" not in line 
                and not line.strip().isdigit() 
                and "WEBVTT" not in line
                and not line.startswith("NOTE")
            )
            # --- NEW: Clean the text after extracting ---
            cleaned_text = clean_transcript_text(full_text)
            transcript_path.write_text(cleaned_text, encoding="utf-8")
            vtt_file.unlink()
            return cleaned_text
    except Exception as e:
        print(f"yt-dlp failed: {e}")
    
    raise Exception("Could not fetch transcript")

def generate_course_overview(full_transcript_text, course_title):
    """Generate 2-4 line overview of the entire course."""
    try:
        model = genai.GenerativeModel(GEMINI_MODEL)
        # Clean text BEFORE sending to AI
        context_text = clean_transcript_text(full_transcript_text)
        context_text = context_text[:5000] if len(context_text) > 5000 else context_text
        
        prompt = f"""Analyze this video course transcript and provide a 2-4 sentence overview of what concepts and topics are taught.
Focus *only* on the educational content. Ignore sponsor messages, introductions, and metadata.
Course Title: {course_title}
Transcript Sample:
{context_text}
Return only the overview, no extra text."""

        response = model.generate_content(prompt, generation_config={"response_mime_type": "text/plain"})
        overview = response.text.strip().replace('\n', ' ')
        return overview
        
    except Exception as e:
        print(f"Course overview generation failed: {e}")
        return f"A comprehensive {course_title} covering essential concepts and practical skills."

def generate_module_summary(module_text, module_num, start_time, end_time):
    """Generate summary for a specific time segment of the video."""
    try:
        model = genai.GenerativeModel(GEMINI_MODEL)
        # Clean text BEFORE sending to AI
        context_text = clean_transcript_text(module_text)
        context_text = context_text[:4000] if len(context_text) > 4000 else context_text
        
        prompt = f"""Summarize the *educational content* taught in this specific segment (minutes {int(start_time/60)}-{int(end_time/60)}) of a video course.
Provide a concise, engaging, 2-3 sentence summary.
IGNORE any sponsor messages, greetings, or off-topic chatter.
Transcript from this time segment:
{context_text}
Return only the summary, no extra text."""

        response = model.generate_content(prompt, generation_config={"response_mime_type": "text/plain"})
        summary = response.text.strip().replace('\n', ' ')
        return summary
        
    except Exception as e:
        print(f"Module summary generation failed: {e}")
        return f"Content covering minutes {int(start_time/60)} to {int(end_time/60)} of the course."

def clean_json_response(text):
    text = re.sub(r'^```json?\s*\n?', '', text, flags=re.MULTILINE)
    text = re.sub(r'\n?```\s*$', '', text, flags=re.MULTILINE)
    return text.strip()

def extract_key_points(module_text):
    """Extract 5 clear, distinct key points from the transcript."""
    try:
        model = genai.GenerativeModel(GEMINI_MODEL)
        # Clean text BEFORE sending to AI
        context_text = clean_transcript_text(module_text)
        context_text = context_text[:4000] if len(context_text) > 4000 else context_text
        
        prompt = f"""Extract exactly 5 key *educational* learning points from this transcript segment.
Ignore metadata, greetings, and sponsor messages. Focus on specific skills or concepts.
Each point should be a single clear concept, 10-20 words maximum.
Transcript:
{context_text}
Return only a JSON array of strings:
["Point 1", "Point 2", "Point 3", "Point 4", "Point 5"]"""

        response = model.generate_content(prompt, generation_config={"response_mime_type": "application/json"})
        key_points = json.loads(response.text)
        
        if isinstance(key_points, list) and len(key_points) > 0:
            return key_points[:5]
        else:
            raise ValueError("Invalid format")
        
    except Exception as e:
        print(f"Key points extraction failed: {e}")
        sentences = [s.strip() for s in module_text.split('.') if 30 < len(s.strip()) < 150]
        return sentences[:5] if sentences else ["Understanding core concepts", "Practical implementation", "Best practices", "Common pitfalls", "Next steps"]

def split_transcript_with_timestamps(transcript_list, daily_study_minutes, video_duration, course_title):
    """Split transcript into modules with time-based AI summaries."""
    VIDEO_TIME_RATIO = 0.85
    QUIZ_TIME_RATIO = 0.15
    
    daily_study_seconds = daily_study_minutes * 60
    video_watch_seconds = int(daily_study_seconds * VIDEO_TIME_RATIO)
    quiz_seconds = int(daily_study_seconds * QUIZ_TIME_RATIO)
    
    num_modules = max(1, int(video_duration / video_watch_seconds))
    seconds_per_module = video_duration / num_modules
    
    print(f"  📊 Time allocation per day:")
    print(f"     Total study time: {daily_study_minutes/60:.2f} hours")
    
    full_text = " ".join(seg['text'] for seg in transcript_list)
    course_overview = generate_course_overview(full_text, course_title)
    print(f"\n✓ Course Overview: {course_overview[:100]}...")
    
    modules = []
    
    for module_idx in range(num_modules):
        start_time = int(module_idx * seconds_per_module)
        end_time = int(min((module_idx + 1) * seconds_per_module, video_duration))
        
        module_segments = [seg for seg in transcript_list if start_time <= seg['start'] < end_time]
        module_text = " ".join(seg['text'] for seg in module_segments)
        
        print(f"\n  📝 Module {module_idx + 1}: {start_time}s - {end_time}s")
        print(f"     Generating time-based summary...")
        summary = generate_module_summary(module_text, module_idx + 1, start_time, end_time)
        print(f"     ✓ {summary[:80]}...")
        
        print(f"     Extracting key points...")
        key_points = extract_key_points(module_text)
        
        module_num = module_idx + 1
        video_duration_hours = round(video_watch_seconds / 3600, 2)
        quiz_duration_hours = round(quiz_seconds / 3600, 2)
        total_duration_hours = round(daily_study_seconds / 3600, 2)
        
        modules.append({
            "day": module_num,
            "title": f"Module {module_num}",
            "description": summary, # This is the AI summary
            "duration": video_duration_hours,
            "quizDuration": quiz_duration_hours,
            "totalDuration": total_duration_hours,
            "motivation": "Stay focused!",
            "completed": False,
            "startTime": start_time,
            "endTime": end_time,
            "module": {
                "description": summary, # This is the AI summary
                "keyPoints": key_points # These are the AI key points
            },
            "quiz": []
        })
    
    return modules, course_overview

# (This function is mostly the same, just added the cleaner)
def split_transcript_without_timestamps(transcript_text, daily_study_minutes, video_duration, course_title):
    """Fallback: Split by word count with AI summaries."""
    VIDEO_TIME_RATIO = 0.85
    QUIZ_TIME_RATIO = 0.15
    
    # --- NEW: Clean the full text first ---
    cleaned_text = clean_transcript_text(transcript_text)
    words = cleaned_text.split()
    total_words = len(words)
    
    daily_study_seconds = daily_study_minutes * 60
    video_watch_seconds = int(daily_study_seconds * VIDEO_TIME_RATIO)
    quiz_seconds = int(daily_study_seconds * QUIZ_TIME_RATIO)
    
    num_modules = max(1, int(video_duration / video_watch_seconds))
    
    print(f"  📊 Time allocation per day: {daily_study_minutes/60:.2f} hours")
    
    words_per_module = total_words // num_modules
    seconds_per_module = video_duration // num_modules
    
    course_overview = generate_course_overview(cleaned_text, course_title)
    print(f"\n✓ Course Overview: {course_overview[:100]}...")
    
    modules = []
    
    for i in range(num_modules):
        start_word = i * words_per_module
        end_word = min((i + 1) * words_per_module, total_words)
        
        module_text = " ".join(words[start_word:end_word])
        module_num = i + 1
        
        start_time = i * seconds_per_module
        end_time = min((i + 1) * seconds_per_module, video_duration)
        
        print(f"\n  📝 Module {module_num}: {start_time}s - {end_time}s")
        print(f"     Generating time-based summary...")
        # Module text is already clean
        summary = generate_module_summary(module_text, module_num, start_time, end_time)
        print(f"     ✓ {summary[:80]}...")
        
        print(f"     Extracting key points...")
        key_points = extract_key_points(module_text)
        
        video_duration_hours = round(video_watch_seconds / 3600, 2)
        quiz_duration_hours = round(quiz_seconds / 3600, 2)
        total_duration_hours = round(daily_study_seconds / 3600, 2)
        
        modules.append({
            "day": module_num,
            "title": f"Module {module_num}",
            "description": summary,
            "duration": video_duration_hours,
            "quizDuration": quiz_duration_hours,
            "totalDuration": total_duration_hours,
            "motivation": "Stay focused!",
            "completed": False,
            "startTime": start_time,
            "endTime": end_time,
            "module": {
                "description": summary,
                "keyPoints": key_points
            },
            "quiz": []
        })
    
    return modules, course_overview

def generate_quiz(module_title, key_points):
    """Generate 5-question quiz using Gemini."""
    try:
        model = genai.GenerativeModel(GEMINI_MODEL)
        key_points_text = "\n".join(f"- {point}" for point in key_points[:5])
        
        prompt = f"""Create exactly 5 multiple choice questions based on {module_title}.
Key points covered:
{key_points_text}
Return ONLY valid JSON:
[
  {{"question": "Clear question?", "options": {{"A": "Option A", "B": "Option B", "C": "Option C", "D": "Option D"}}, "correct_answer": "A"}}
]"""

        response = model.generate_content(prompt, generation_config={"response_mime_type": "application/json"})
        quiz = json.loads(response.text)
        return quiz[:5] if isinstance(quiz, list) else []
        
    except Exception as e:
        print(f"Quiz generation failed: {e}")
        return [{
            "question": f"What is a key concept from {module_title}?",
            "options": {
                "A": key_points[0] if len(key_points) > 0 else "Concept A",
                "B": key_points[1] if len(key_points) > 1 else "Concept B",
                "C": key_points[2] if len(key_points) > 2 else "Concept C",
                "D": "None of the above"
            },
            "correct_answer": "A"
        }]

def generate_coding_challenge(module_title, key_points):
    """Generate a single coding challenge for the module."""
    try:
        model = genai.GenerativeModel(GEMINI_MODEL)
        key_points_text = "\n".join(f"- {point}" for point in key_points)
        
        prompt = f"""
You are an expert curriculum developer. Based on the key points from a video module, create a *single* simple coding challenge.
The module is: "{module_title}"
The key points are:
{key_points_text}
Analyze the key points. If they are clearly about web development (HTML, CSS, JavaScript), create a challenge for that.
If the topics are generic (e.g., "data structures", "loops", "variables"), create a simple JavaScript challenge.
If the topics are *not* coding-related at all (e.g., "History of Rome", "Photoshop basics"), return a 'none' type.
Return ONLY a valid JSON object with this exact structure:
{{
  "type": "html" or "javascript" or "none",
  "title": "A short, engaging challenge title",
  "question": "A 2-4 sentence description of the challenge. What does the user need to do?",
  "starting_code": {{
    "html": "...",
    "css": "...",
    "js": "..."
  }},
  "solution": {{
    "html": "...",
    "css": "...",
    "js": "..."
  }}
}}
"""
        response = model.generate_content(prompt, generation_config={"response_mime_type": "application/json"})
        challenge = json.loads(response.text)
        return challenge
        
    except Exception as e:
        print(f"Coding challenge generation failed: {e}")
        return {"type": "none", "title": "", "question": "", "starting_code": {"html": "", "css": "", "js": ""}, "solution": {"html": "", "css": "", "js": ""}}

def get_ai_hint(challenge_question, user_code, solution, try_count):
    """Generate a hint for the user's code."""
    try:
        model = genai.GenerativeModel(GEMINI_MODEL)
        user_code_text = f"HTML:\n{user_code['html']}\n\nCSS:\n{user_code['css']}\n\nJS:\n{user_code['js']}"
        solution_text = f"HTML:\n{solution['html']}\n\nCSS:\n{solution['css']}\n\nJS:\n{solution['js']}"
        
        if try_count >= 3:
            # Give the full answer
            prompt = f"""
The user is stuck on a coding challenge and has asked for help {try_count} times. Give them the full solution and a brief, friendly explanation of how it works.
The challenge was: {challenge_question}
The correct solution is:
{solution_text}
"""
        else:
            # Give a hint
            prompt = f"""
You are a helpful Socratic coding tutor. A student is stuck on a challenge. Do NOT give them the direct answer.
Instead, give them a 1-2 sentence hint to guide them.
The challenge was: {challenge_question}
The user's code is:
{user_code_text}
Give a small, targeted hint. For example: "You're close! Take another look at your CSS `background-color` property." or "Remember how to declare a function in JavaScript? Check your syntax on line 1."
"""
        response = model.generate_content(prompt, generation_config={"response_mime_type": "text/plain"})
        hint = response.text.strip()
        
        # Add a flag to tell the frontend if the solution was given
        solution_given = (try_count >= 3)
        return {"hint": hint, "solution_given": solution_given}
        
    except Exception as e:
        print(f"AI hint generation failed: {e}")
        return {"hint": "I'm having trouble thinking of a hint right now. Why don't you try checking your console for errors?", "solution_given": False}


# --- Routes ---

@app.route("/generate-plan", methods=["POST", "OPTIONS"])
def generate_plan():
    if request.method == "OPTIONS":
        return jsonify({"status": "ok"})
        
    try:
        data = request.get_json()
        course_title = data.get("courseTitle", "Course")
        course_link = data.get("courseLink")
        daily_hours = float(data.get("dailyStudyHours", 1))
        
        video_id = extract_youtube_id(course_link)
        if not video_id:
            return jsonify({"error": "Invalid YouTube link"}), 400

        print(f"\n{'='*60}")
        print(f"📚 Processing: {course_title}")
        print(f"🎥 Video ID: {video_id}")
        
        video_duration = get_video_duration(video_id)
        print(f"✓ Video duration: {video_duration}s ({video_duration/60:.1f} minutes)")
        
        transcript_with_timestamps = get_transcript_with_timestamps(video_id)
        
        if transcript_with_timestamps:
            print(f"✓ Got transcript with timestamps: {len(transcript_with_timestamps)} segments")
            daily_plan, course_overview = split_transcript_with_timestamps(
                transcript_with_timestamps, 
                daily_hours * 60,
                video_duration,
                course_title
            )
        else:
            print("⚠ Using word-based splitting")
            transcript_text = get_transcript_text(video_id)
            print(f"✓ Got transcript: {len(transcript_text.split())} words")
            daily_plan, course_overview = split_transcript_without_timestamps(
                transcript_text,
                daily_hours * 60,
                video_duration,
                course_title
            )

        print(f"\n✓ Created {len(daily_plan)} modules\n")

        response_data = {
            "courseTitle": course_title,
            "courseDescription": course_overview,
            "videoID": video_id,
            "dailyPlan": daily_plan,
            "streak": 0,
            "progress": 0
        }
        
        print(f"✅ Plan generated successfully!\n")
        return jsonify(response_data)
        
    except Exception as e:
        print(f"\n❌ Error: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@app.route("/generate-quiz", methods=["POST", "OPTIONS"])
def generate_quiz_route():
    if request.method == "OPTIONS":
        return jsonify({"status": "ok"})
    try:
        data = request.get_json()
        module_title = data.get("moduleTitle")
        key_points = data.get("keyPoints", [])
        
        print(f"📝 Generating quiz for: {module_title}")
        quiz = generate_quiz(module_title, key_points)
        print(f"✓ Generated {len(quiz)} questions")
        
        return jsonify(quiz)
    except Exception as e:
        print(f"Error: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route("/generate-challenge", methods=["POST", "OPTIONS"])
def generate_challenge_route():
    if request.method == "OPTIONS":
        return jsonify({"status": "ok"})
    try:
        data = request.get_json()
        module_title = data.get("moduleTitle")
        key_points = data.get("keyPoints", [])
        
        print(f"🚀 Generating challenge for: {module_title}")
        challenge = generate_coding_challenge(module_title, key_points)
        print(f"✓ Generated challenge, type: {challenge.get('type')}")
        
        return jsonify(challenge)
    except Exception as e:
        print(f"Error: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route("/get-hint", methods=["POST", "OPTIONS"])
def get_hint_route():
    if request.method == "OPTIONS":
        return jsonify({"status": "ok"})
    try:
        data = request.get_json()
        challenge_question = data.get("challenge_question")
        user_code = data.get("user_code")
        solution = data.get("solution")
        try_count = int(data.get("try_count", 1))
        
        print(f"🤖 Generating hint for user, attempt #{try_count}")
        hint_data = get_ai_hint(challenge_question, user_code, solution, try_count)
        print(f"✓ Hint generated. Solution given: {hint_data.get('solution_given')}")
        
        return jsonify(hint_data)
    except Exception as e:
        print(f"Error: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route("/get-motivation", methods=["POST", "OPTIONS"])
def get_motivation():
    if request.method == "OPTIONS":
        return jsonify({"status": "ok"})
    try:
        data = request.get_json()
        module_title = data.get("moduleTitle")
        course_title = data.get("courseTitle")
        message = f"🎉 Excellent work! You've completed {module_title} in {course_title}. Keep building your knowledge!"
        return jsonify({"message": message})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# --- NEW ROUTE: Generate .docx Notes ---
@app.route("/generate-notes-doc", methods=["POST", "OPTIONS"])
def generate_notes_doc():
    if request.method == "OPTIONS":
        return jsonify({"status": "ok"})
    try:
        data = request.get_json()
        module_title = data.get("moduleTitle", "Module Notes")
        summary = data.get("summary", "No summary provided.")
        key_points = data.get("keyPoints", [])

        print(f"📄 Generating .docx for: {module_title}")

        # Create a new Document
        document = Document()
        document.add_heading(module_title, level=1)

        # Add Summary
        document.add_heading('Module Summary', level=2)
        document.add_paragraph(summary)

        # Add Key Points
        document.add_heading('Key Learning Points', level=2)
        for point in key_points:
            document.add_paragraph(point, style='List Bullet')

        # Save to a memory buffer
        file_stream = io.BytesIO()
        document.save(file_stream)
        file_stream.seek(0) # Rewind the buffer

        print("✓ Document created successfully.")

        return jsonify({
            "file_blob": file_stream.read().hex(), # Send as hex string to avoid JSON issues
            "file_name": f"{module_title}_Notes.docx"
        })
        
    except Exception as e:
        print(f"Error generating .docx: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route("/health", methods=["GET"])
def health():
    return jsonify({"status": "healthy", "message": "StudyMate API is running"})

if __name__ == "__main__":
    print("\n" + "="*60)
    print("🎓 StudyMate API Server")
    print("="*60)
    print(f"Gemini Model: {GEMINI_MODEL}")
    print(f"Server: http://127.0.0.1:5000")
    print("="*60 + "\n")
    app.run(debug=True, host='127.0.0.1', port=5000)